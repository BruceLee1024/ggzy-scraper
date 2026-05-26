#!/usr/bin/env python3
"""
JSON 转 Excel - 全国公共资源交易平台招标公告（含详情信息）
支持基础字段 + 详情字段（金额、截止时间、资金来源等）
"""
import json, openpyxl, sys, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def json_to_excel(json_path, excel_path, detail_fields=True):
    """
    将抓取的JSON数据转为格式化Excel
    
    detail_fields=True时包含详情页字段:
      预算(万元), 保证金(万元), 投标截止, 文件获取截止,
      资金来源, 建设规模, 计划工期, 招标代理
    """
    with open(json_path, "r", encoding="utf-8") as f:
        records = json.load(f)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "招标公告"
    
    # 基础字段
    headers = ["序号", "标题", "发布时间", "省份", "来源平台",
               "业务类型", "信息类型", "行业", "URL"]
    
    # 详情字段
    detail_headers = []
    if detail_fields:
        detail_headers = [
            "预算(万元)", "保证金(万元)", "投标截止时间",
            "文件获取截止", "资金来源", "建设规模",
            "计划工期(天)", "招标代理"
        ]
    
    all_headers = headers + detail_headers
    
    # 样式
    hfont = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    dfont = Font(size=10)
    lfont = Font(size=10, color="0563C1", underline="single")
    mfont = Font(size=10, color="C00000")  # 金额红色
    border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )
    altfill = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
    budget_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    center_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # 写入表头
    for col, h in enumerate(all_headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.fill, c.alignment, c.border = hfont, hfill, center, border
    
    # 列宽
    widths = [6, 40, 12, 8, 24, 10, 12, 12, 60]
    if detail_fields:
        widths += [10, 10, 16, 14, 20, 50, 10, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    
    # 写入数据
    for idx, rec in enumerate(records, 1):
        row = idx + 1
        url_path = rec.get("url", "")
        full_url = f"https://www.ggzy.gov.cn{url_path}" if url_path and not url_path.startswith("http") else url_path
        
        # 基础值
        vals = [
            idx, rec.get("title", ""), rec.get("publishTime", ""),
            rec.get("provinceText", ""), rec.get("transactionSourcesPlatformText", ""),
            rec.get("businessTypeText", ""), rec.get("informationTypeText", ""),
            rec.get("industryTypeText", ""), full_url
        ]
        
        # 详情值
        if detail_fields:
            budget = rec.get("budget_money", "")
            deposit = rec.get("deposit_money", "")
            vals += [
                budget, deposit,
                rec.get("bid_deadline", ""), rec.get("doc_deadline", ""),
                rec.get("funding_source", ""), rec.get("scale", ""),
                rec.get("plan_days", ""), rec.get("tender_agent", "")
            ]
        
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font, c.border = dfont, border
            if col in (1, 3, 4, 6, 7, 8):
                c.alignment = center
            elif col in (10, 11, 12, 13, 16):  # 金额和数字列
                c.alignment = center_wrap
            else:
                c.alignment = left
            
            # 金额高亮
            if col in (10, 11) and val:
                c.font = mfont
                c.fill = budget_fill
            
            # 隔行色
            if idx % 2 == 0 and col not in (10, 11):
                c.fill = altfill
        
        # URL 超链接
        uc = ws.cell(row=row, column=9)
        if full_url:
            uc.hyperlink, uc.font = full_url, lfont
    
    ws.freeze_panes = "A2"
    last_col = openpyxl.utils.get_column_letter(len(all_headers))
    ws.auto_filter.ref = f"A1:{last_col}{len(records)+1}"
    
    wb.save(excel_path)
    print(f"Saved {len(records)} records -> {excel_path}")
    if detail_fields:
        has_money = sum(1 for r in records if r.get("budget_money"))
        print(f"  含预算金额: {has_money}/{len(records)} 条")
    return len(records)


if __name__ == "__main__":
    json_path = sys.argv[1] if len(sys.argv) > 1 else "/tmp/ggzy_records.json"
    excel_path = sys.argv[2] if len(sys.argv) > 2 else os.path.expanduser("~/Desktop/招标公告.xlsx")
    mode = sys.argv[3] if len(sys.argv) > 3 else "detail"
    detail_fields = (mode.lower() != "basic")
    json_to_excel(json_path, excel_path, detail_fields=detail_fields)
